# -*- coding: utf-8 -*-
"""
Erstellt die PowerPoint-Präsentation zur Abgabe 3 (Musikgenerierung mit RNNs).

Voraussetzung: Das Notebook `ubung03.ipynb` wurde einmal komplett ausgeführt,
damit folgende Dateien existieren:
  - ergebnisse.json       (alle Metriken & Konfiguration)
  - beispiele.txt         (generierte Musikstücke)
  - kurven_loss.png, kurven_accuracy.png, vergleich_lstm_gru.png

Aufruf:  python erstelle_praesentation.py
Ausgabe: praesentation_ubung03.pptx
"""
import json
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# ----------------------------------------------------------------- Farben & Layout
NAVY      = RGBColor(0x16, 0x32, 0x4F)   # dunkles Blau für Titel-Hintergrund
BLAU      = RGBColor(0x2A, 0x78, 0xD6)   # Akzent 1 (gleiche Farbe wie in den Plots)
GRUEN     = RGBColor(0x00, 0x83, 0x00)   # Akzent 2
TINTE     = RGBColor(0x0B, 0x0B, 0x0B)   # Haupttext
GRAU      = RGBColor(0x52, 0x51, 0x4E)   # Nebentext
HELLGRAU  = RGBColor(0xF0, 0xEF, 0xEC)   # Hintergrund für Code-/Zahlen-Kacheln
WEISS     = RGBColor(0xFF, 0xFF, 0xFF)

FONT      = "Segoe UI"
MONO      = "Consolas"

BREITE = Inches(13.333)   # 16:9
HOEHE  = Inches(7.5)


# ----------------------------------------------------------------- kleine Helfer
def neue_folie(prs):
    """Leere Folie (Layout 6 = 'blank') hinzufügen."""
    return prs.slides.add_slide(prs.slide_layouts[6])


def rechteck(folie, x, y, b, h, farbe, rund=False):
    """Gefülltes Rechteck ohne Rand und ohne Schatten."""
    form_typ = MSO_SHAPE.ROUNDED_RECTANGLE if rund else MSO_SHAPE.RECTANGLE
    form = folie.shapes.add_shape(form_typ, x, y, b, h)
    form.fill.solid()
    form.fill.fore_color.rgb = farbe
    form.line.fill.background()
    form.shadow.inherit = False
    return form


def textfeld(folie, x, y, b, h, text, groesse=16, farbe=TINTE, fett=False,
             font=FONT, ausrichtung=PP_ALIGN.LEFT, anker=MSO_ANCHOR.TOP,
             abstand_nach=4):
    """Mehrzeiliges Textfeld; jede Zeile des Strings wird ein eigener Absatz."""
    box = folie.shapes.add_textbox(x, y, b, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anker
    for i, zeile in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ausrichtung
        p.space_after = Pt(abstand_nach)
        run = p.add_run()
        run.text = zeile
        run.font.name = font
        run.font.size = Pt(groesse)
        run.font.bold = fett
        run.font.color.rgb = farbe
    return box


def inhaltsfolie(prs, titel):
    """Standard-Folie: Titel oben + blaue Akzentlinie darunter."""
    folie = neue_folie(prs)
    textfeld(folie, Inches(0.55), Inches(0.32), Inches(12.2), Inches(0.75),
             titel, groesse=30, fett=True)
    rechteck(folie, Inches(0.6), Inches(1.08), Inches(1.1), Pt(4), BLAU)
    return folie


def code_kachel(folie, x, y, b, h, code_text, groesse=11):
    """Hellgraue Kachel mit Monospace-Text (für ABC-Notation)."""
    rechteck(folie, x, y, b, h, HELLGRAU, rund=True)
    textfeld(folie, x + Inches(0.25), y + Inches(0.18), b - Inches(0.5),
             h - Inches(0.36), code_text, groesse=groesse, font=MONO,
             farbe=TINTE, abstand_nach=0)


def stat_kachel(folie, x, y, b, h, wert, beschriftung, farbe=BLAU):
    """Kachel mit großer Zahl + Beschriftung (Hero-Zahl)."""
    rechteck(folie, x, y, b, h, HELLGRAU, rund=True)
    textfeld(folie, x, y + Inches(0.35), b, Inches(1.3), wert, groesse=54,
             fett=True, farbe=farbe, ausrichtung=PP_ALIGN.CENTER)
    textfeld(folie, x, y + h - Inches(0.85), b, Inches(0.6), beschriftung,
             groesse=15, farbe=GRAU, ausrichtung=PP_ALIGN.CENTER)


def architektur_box(folie, x, y, b, h, titel, untertitel, farbe):
    """Ein Baustein des Architektur-Diagramms."""
    kasten = rechteck(folie, x, y, b, h, farbe, rund=True)
    tf = kasten.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, (zeile, gr, fett) in enumerate([(titel, 15, True), (untertitel, 11, False)]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = zeile
        run.font.name = FONT
        run.font.size = Pt(gr)
        run.font.bold = fett
        run.font.color.rgb = WEISS


def pfeil(folie, x, y, laenge=Inches(0.42)):
    """Kleiner Pfeil zwischen zwei Architektur-Boxen."""
    form = folie.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y, laenge, Inches(0.28))
    form.fill.solid()
    form.fill.fore_color.rgb = GRAU
    form.line.fill.background()
    form.shadow.inherit = False


def bild_oder_hinweis(folie, pfad, x, y, breite):
    """Bild einfügen; falls es fehlt, einen Hinweis anzeigen."""
    if os.path.exists(pfad):
        folie.shapes.add_picture(pfad, x, y, width=breite)
    else:
        textfeld(folie, x, y, breite, Inches(1), f"[Bild fehlt: {pfad}]",
                 groesse=14, farbe=GRAU)


# ----------------------------------------------------------------- Daten einlesen
if not os.path.exists("ergebnisse.json"):
    sys.exit("FEHLER: ergebnisse.json nicht gefunden.\n"
             "Bitte zuerst das Notebook ubung03.ipynb komplett ausführen.")

with open("ergebnisse.json", encoding="utf-8") as f:
    erg = json.load(f)

cfg = erg["config"]
lstm, gru = erg["lstm"], erg["gru"]

if os.path.exists("beispiele.txt"):
    with open("beispiele.txt", encoding="utf-8") as f:
        generierte_stuecke = f.read().split("\n\n")
else:
    generierte_stuecke = ["(beispiele.txt fehlt)"]


def kuerze(text, max_zeilen=14, max_breite=52):
    """Text auf eine Kachel-freundliche Größe kürzen."""
    zeilen = [z[:max_breite] for z in text.split("\n")[:max_zeilen]]
    return "\n".join(zeilen)


# ----------------------------------------------------------------- Präsentation
prs = Presentation()
prs.slide_width = BREITE
prs.slide_height = HOEHE

# --- Folie 1: Titel -----------------------------------------------------------
folie = neue_folie(prs)
rechteck(folie, 0, 0, BREITE, HOEHE, NAVY)
rechteck(folie, Inches(0.9), Inches(2.1), Inches(1.4), Pt(5), GRUEN)
textfeld(folie, Inches(0.85), Inches(2.35), Inches(11.6), Inches(1.8),
         "Musikgenerierung mit RNNs", groesse=48, fett=True, farbe=WEISS)
textfeld(folie, Inches(0.9), Inches(3.55), Inches(11.5), Inches(0.6),
         "Ein LSTM lernt irische Volksmusik in ABC-Notation – und komponiert selbst.",
         groesse=19, farbe=RGBColor(0xC9, 0xD6, 0xE4))
textfeld(folie, Inches(0.9), Inches(5.9), Inches(11.5), Inches(1.0),
         "Einführung in Deep Learning · Abgabe 3\nJordan Pokem · Leslie Tsafack",
         groesse=15, farbe=RGBColor(0xC9, 0xD6, 0xE4))

# --- Folie 2: Aufgabe ----------------------------------------------------------
folie = inhaltsfolie(prs, "Die Aufgabe")
textfeld(folie, Inches(0.6), Inches(1.5), Inches(6.6), Inches(5.2),
         "•  Ziel: neue irische Melodien automatisch generieren\n"
         "•  Musik liegt als Text vor (ABC-Notation) →\n"
         "    Musikgenerierung = Sprachmodellierung\n"
         "•  Ein RNN/LSTM lernt, das jeweils nächste Zeichen\n"
         "    vorherzusagen (Vorlesung 8 & 9)\n"
         "•  Evaluation mit Top-1- und Top-5-Accuracy\n"
         "•  Training komplett mit wandb visualisiert",
         groesse=18)
# rechte Seite: Mini-Pipeline von oben nach unten
stationen = [("ABC-Text", "irische Volkslieder (IrishMAN)"),
             ("Zeichen-Tokenizer", "Text → Zahlen-IDs"),
             ("Embedding + LSTM", "lernt das nächste Zeichen"),
             ("Sampling", "Zeichen für Zeichen → neue Melodie")]
y = Inches(1.45)
for i, (t, u) in enumerate(stationen):
    architektur_box(folie, Inches(8.1), y, Inches(4.4), Inches(0.95), t, u,
                    BLAU if i % 2 == 0 else GRUEN)
    y += Inches(1.32)

# --- Folie 3: Datensatz --------------------------------------------------------
folie = inhaltsfolie(prs, "Datensatz: IrishMAN")
textfeld(folie, Inches(0.6), Inches(1.5), Inches(6.2), Inches(5.2),
         "•  Irish Massive ABC Notation (Hugging Face)\n"
         f"•  Verwendet: {cfg['num_tunes']:,} Stücke aus dem Train-Split\n"
         f"•  Vokabular: {erg['vocab_groesse']} verschiedene Zeichen\n"
         f"•  {erg['anzahl_sequenzen']:,} Trainingssequenzen "
         f"(Fensterlänge {cfg['seq_len']})\n"
         "•  Split: 80 % Train / 10 % Validation / 10 % Test\n\n"
         "ABC-Notation:\n"
         "•  Kopfzeilen: M: Taktart, L: Notenlänge, K: Tonart\n"
         "•  Noten A–G / a–g, Taktstriche |, Pausen z",
         groesse=17)
code_kachel(folie, Inches(7.3), Inches(1.4), Inches(5.4), Inches(5.4),
            kuerze(erg.get("beispiel_original", ""), max_zeilen=16), groesse=12)
textfeld(folie, Inches(7.3), Inches(6.85), Inches(5.4), Inches(0.4),
         "Original-Stück aus dem Datensatz", groesse=12, farbe=GRAU,
         ausrichtung=PP_ALIGN.CENTER)

# --- Folie 4: Vorverarbeitung --------------------------------------------------
folie = inhaltsfolie(prs, "Vorverarbeitung: Tokenizer & gleitendes Fenster")
textfeld(folie, Inches(0.6), Inches(1.5), Inches(6.4), Inches(4.5),
         "•  Zeichen-Tokenizer: jedes Zeichen = 1 Token\n"
         "    (einfach + kein OOV-Problem, Vorlesung 9)\n"
         "•  encode: Text → IDs  ·  decode: IDs → Text\n"
         f"•  Gleitendes Fenster der Länge {cfg['seq_len']},\n"
         f"    Schrittweite {cfg['schritt']}\n"
         "•  Ziel-Sequenz = Input um 1 Zeichen verschoben\n"
         "    → an jeder Position wird das nächste Zeichen gelernt",
         groesse=18)
code_kachel(folie, Inches(7.4), Inches(2.0), Inches(5.2), Inches(1.0),
            'Input:  "K:D\\n|:A2B c2d|"', groesse=15)
pfeil_form = folie.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(9.75),
                                    Inches(3.15), Inches(0.5), Inches(0.55))
pfeil_form.fill.solid()
pfeil_form.fill.fore_color.rgb = GRAU
pfeil_form.line.fill.background()
pfeil_form.shadow.inherit = False
code_kachel(folie, Inches(7.4), Inches(3.85), Inches(5.2), Inches(1.0),
            'Ziel:   ":D\\n|:A2B c2d|e"', groesse=15)
textfeld(folie, Inches(7.4), Inches(5.0), Inches(5.2), Inches(0.5),
         "um 1 Zeichen verschoben", groesse=13, farbe=GRAU,
         ausrichtung=PP_ALIGN.CENTER)

# --- Folie 5: Architektur ------------------------------------------------------
folie = inhaltsfolie(prs, "Architektur: Embedding → LSTM → Linear")
boxen = [("Zeichen-IDs", f"Batch × {cfg['seq_len']}", GRAU),
         ("Embedding", f"Dim {cfg['embed_dim']}", BLAU),
         ("LSTM", f"{cfg['num_layers']} × {cfg['hidden_size']}", GRUEN),
         ("Linear", f"→ {erg['vocab_groesse']} Logits", BLAU),
         ("Softmax", "nächstes Zeichen", GRAU)]
x = Inches(0.55)
for i, (t, u, farbe) in enumerate(boxen):
    architektur_box(folie, x, Inches(2.6), Inches(2.15), Inches(1.35), t, u, farbe)
    x += Inches(2.15)
    if i < len(boxen) - 1:
        pfeil(folie, x + Inches(0.04), Inches(3.15))
        x += Inches(0.5)
textfeld(folie, Inches(0.6), Inches(4.6), Inches(12.2), Inches(1.6),
         f"•  {erg['anzahl_parameter']:,} lernbare Parameter\n"
         "•  Der Hidden State trägt den musikalischen Kontext von Zeitschritt zu Zeitschritt\n"
         "•  Dieselbe Klasse kann auch GRU oder einfaches RNN verwenden (→ Ablation)",
         groesse=17)

# --- Folie 6: Training ---------------------------------------------------------
folie = inhaltsfolie(prs, "Training")
zeilen = [("Batch Size", str(cfg["batch_size"])),
          ("Learning Rate", str(cfg["learning_rate"])),
          ("Epochen", str(cfg["num_epochs"])),
          ("Hidden Size", str(cfg["hidden_size"])),
          ("RNN-Schichten", str(cfg["num_layers"])),
          ("Sequenzlänge", str(cfg["seq_len"]))]
tabelle = folie.shapes.add_table(len(zeilen) + 1, 2, Inches(0.6), Inches(1.5),
                                 Inches(5.2), Inches(4.4)).table
tabelle.cell(0, 0).text = "Hyperparameter"
tabelle.cell(0, 1).text = "Wert"
for r, (name, wert) in enumerate(zeilen, start=1):
    tabelle.cell(r, 0).text = name
    tabelle.cell(r, 1).text = wert
for r in range(len(zeilen) + 1):
    for c in range(2):
        for p in tabelle.cell(r, c).text_frame.paragraphs:
            for run in p.runs:
                run.font.name = FONT
                run.font.size = Pt(15)
                run.font.bold = (r == 0)
textfeld(folie, Inches(6.6), Inches(1.5), Inches(6.2), Inches(5.0),
         "•  Loss: CrossEntropy über alle Zeitschritte\n"
         "•  Optimizer: Adam\n"
         "•  Metriken pro Epoche (Train + Validation):\n"
         "    Loss, Top-1- und Top-5-Accuracy\n\n"
         "•  Alles live in wandb geloggt:\n"
         "    Projekt: eidl-thm /\n"
         "    4.block_Jordan_Pokem_Leslie_Tsafack_RNN\n"
         "•  Ein Run pro Experiment (LSTM, GRU)",
         groesse=18)

# --- Folie 7: Trainingskurven --------------------------------------------------
folie = inhaltsfolie(prs, "Trainingskurven (LSTM)")
bild_oder_hinweis(folie, "kurven_loss.png", Inches(0.55), Inches(1.7), Inches(6.0))
bild_oder_hinweis(folie, "kurven_accuracy.png", Inches(6.85), Inches(1.7), Inches(6.0))
textfeld(folie, Inches(0.6), Inches(5.6), Inches(12.2), Inches(1.2),
         "•  Loss fällt stetig, Train und Validation bleiben nah beieinander (kein starkes Overfitting)\n"
         "•  Interaktive Kurven und alle Runs: wandb.ai → eidl-thm",
         groesse=16)

# --- Folie 8: Evaluation -------------------------------------------------------
folie = inhaltsfolie(prs, "Evaluation auf dem Test-Set")
stat_kachel(folie, Inches(1.3), Inches(1.9), Inches(4.6), Inches(2.6),
            f"{lstm['test']['top1']:.1%}", "Top-1 Accuracy (LSTM)", BLAU)
stat_kachel(folie, Inches(7.3), Inches(1.9), Inches(4.6), Inches(2.6),
            f"{lstm['test']['top5']:.1%}", "Top-5 Accuracy (LSTM)", GRUEN)
zufall = 1.0 / max(erg["vocab_groesse"], 1)
textfeld(folie, Inches(0.6), Inches(5.1), Inches(12.2), Inches(1.6),
         "•  Definition (Übungsblatt): Anteil der Test-Sequenzen, bei denen das nächste Zeichen\n"
         "    richtig (Top-1) bzw. unter den 5 wahrscheinlichsten (Top-5) vorhergesagt wird\n"
         f"•  Zum Vergleich: zufälliges Raten läge bei ≈ {zufall:.1%} "
         f"(Vokabular mit {erg['vocab_groesse']} Zeichen)",
         groesse=16)

# --- Folie 9: Generierte Musik -------------------------------------------------
folie = inhaltsfolie(prs, "Generierte Musik")
code_kachel(folie, Inches(0.6), Inches(1.5), Inches(6.4), Inches(5.3),
            kuerze(generierte_stuecke[0], max_zeilen=15, max_breite=48), groesse=13)
textfeld(folie, Inches(7.5), Inches(1.6), Inches(5.2), Inches(5.0),
         "•  Seed = typischer ABC-Anfang\n    (X: / L: / M: / K:)\n"
         "•  Autoregressives Sampling:\n"
         "    Softmax → Stichprobe →\n"
         "    Zeichen anhängen → weiter\n"
         "•  Stichprobe statt Argmax:\n"
         "    jede Melodie wird anders,\n"
         "    keine Endlos-Wiederholungen\n\n"
         "•  Demo: generierte_musik.wav 🔊",
         groesse=18)

# --- Folie 10: Bonus -----------------------------------------------------------
folie = inhaltsfolie(prs, "Bonus: Ablation, Validity Check & Audio")
bild_oder_hinweis(folie, "vergleich_lstm_gru.png", Inches(0.55), Inches(1.6), Inches(5.9))
textfeld(folie, Inches(6.8), Inches(1.5), Inches(6.0), Inches(5.4),
         "Ablation LSTM vs. GRU (2 P.)\n"
         f"•  LSTM: Top-1 {lstm['test']['top1']:.1%} · Top-5 {lstm['test']['top5']:.1%}\n"
         f"•  GRU:  Top-1 {gru['test']['top1']:.1%} · Top-5 {gru['test']['top5']:.1%}\n\n"
         "Validity Check (2 P.)\n"
         "•  6 Grammatik-Regeln der ABC-Notation\n"
         "    (Kopfzeilen, Tonart, Taktstriche, Noten, Schluss)\n"
         f"•  Score der generierten Stücke: {erg['validity_score']:.0%}\n\n"
         "Kreativität (1 P.)\n"
         "•  Eigener Mini-ABC-Parser + Sinus-Synthese\n"
         "•  Generierte Melodie als hörbare WAV-Datei",
         groesse=16)

# --- Folie 11: Herausforderungen -----------------------------------------------
folie = inhaltsfolie(prs, "Herausforderungen & Lösungen")
textfeld(folie, Inches(0.6), Inches(1.6), Inches(12.2), Inches(5.4),
         "•  Datenmenge vs. Trainingszeit  →  Teil-Set + Schrittweite des Fensters als Stellschrauben\n\n"
         "•  Metrik-Definition vom Übungsblatt (pro Sequenz)  →  Accuracy nur auf dem letzten\n"
         "    Zeitschritt berechnet, Loss aber über alle Zeitschritte (schnelleres Lernen)\n\n"
         "•  Repetitive Musik bei Argmax  →  Sampling aus der Softmax-Verteilung (torch.multinomial)\n\n"
         "•  Lange Stücke verlieren Struktur  →  Kontext des Modells ist auf die Fensterlänge begrenzt;\n"
         "    Stücke werden an der ersten Leerzeile abgeschnitten\n\n"
         "•  ABC-Syntax bewerten  →  eigener Regel-Check statt nur Accuracy",
         groesse=17)

# --- Folie 12: Fazit -----------------------------------------------------------
folie = inhaltsfolie(prs, "Fazit & Ausblick")
textfeld(folie, Inches(0.6), Inches(1.6), Inches(12.2), Inches(4.2),
         "•  Ein einfaches LSTM-Sprachmodell reicht, um plausible irische Melodien zu erzeugen\n"
         f"•  Test-Set: Top-1 {lstm['test']['top1']:.1%} · Top-5 {lstm['test']['top5']:.1%} "
         f"(Zufall ≈ {zufall:.1%})\n"
         "•  LSTM und GRU liegen nahezu gleichauf – GRU trainiert mit weniger Parametern\n"
         "•  Ausblick: längerer Kontext, größeres Modell, Subword-Tokenizer (BPE),\n"
         "    Attention/Transformer (Vorlesung 10)",
         groesse=18)
textfeld(folie, Inches(0.6), Inches(6.0), Inches(12.2), Inches(0.8),
         "Vielen Dank! – Fragen?", groesse=24, fett=True, farbe=BLAU)

# ----------------------------------------------------------------- speichern
ziel = "praesentation_ubung03.pptx"
prs.save(ziel)
print(f"Präsentation gespeichert: {ziel} ({len(prs.slides._sldIdLst)} Folien)")
