# -*- coding: utf-8 -*-
"""
Fügt der Präsentation `praesentation_ubung03.pptx` die Referentennotizen hinzu
(sichtbar nur in der Referentenansicht, nicht auf dem Beamer).

Das Skript ändert NUR die Notizen – die Folien selbst bleiben unangetastet.
Es ist idempotent: mehrmaliges Ausführen überschreibt die Notizen einfach neu.

Aufruf (aus dem Ordner util/):  python notizen_hinzufuegen.py
"""
import os
import shutil
import sys

from pptx import Presentation

PPTX = "praesentation_ubung03.pptx"

# ----------------------------------------------------------------- Notizen
# Schlüssel = Foliennummer (1-basiert, wie in PowerPoint angezeigt)
NOTIZEN = {

1: """▶ BEGRÜSSUNG
"Guten Morgen. Wir sind Jordan Pokem und Leslie Tsafack.
Unser Thema ist die Abgabe 3: Musikgenerierung mit RNNs.
Wir haben ein LSTM trainiert, das irische Volksmusik in ABC-Notation
liest und daraus selbst neue Melodien komponiert.
Am Ende spielen wir Ihnen ein Stück vor, das das Modell erfunden hat."

▶ ABLAUF ANKÜNDIGEN
Aufgabe → Daten → Architektur → Training → Ergebnisse → Bonus → Fazit.

💡 TIPP: ruhig atmen, Blickkontakt, dann erst weiterklicken.""",

2: """▶ KERNIDEE (der wichtigste Satz der ganzen Präsentation!)
"Musik in ABC-Notation ist einfach Text.
Damit wird Musikgenerierung zu einem Sprachmodellierungs-Problem:
Das Netz lernt, das jeweils nächste Zeichen vorherzusagen."

▶ WARUM EIN RNN UND KEIN NORMALES NETZ?
"Ein klassisches Netz erwartet eine feste Eingabegröße und erfasst die
zeitliche Reihenfolge nicht. Musik ist aber eine Sequenz: variable Länge,
und die Reihenfolge ist entscheidend – dieselben Noten in anderer
Reihenfolge ergeben eine völlig andere Melodie."

▶ RECHTS: die Pipeline von oben nach unten kurz durchgehen
ABC-Text → Zeichen-Tokenizer → Embedding + LSTM → Sampling → neue Melodie.""",

3: """▶ SAGEN
"Wir verwenden den IrishMAN-Datensatz von Hugging Face – irische
Volkslieder in ABC-Notation."

📊 ZAHLEN
• Datensatz gesamt: 214.122 Stücke
• Von uns genutzt: 8.000 Stücke  (bewusster Kompromiss – siehe Folie 11)
• Vokabular: 92 verschiedene Zeichen
• 46.686 Trainingssequenzen
• Split: 80 % Train / 10 % Validation / 10 % Test

▶ ABC-NOTATION KURZ ERKLÄREN (rechts auf das Beispiel zeigen)
• Kopfzeilen: M: = Taktart, L: = Grundnotenlänge, K: = Tonart
• Noten: A–G (tiefe Oktave) und a–g (hohe Oktave)
• | = Taktstrich, z = Pause

❓ FALLS GEFRAGT "Warum nur 8.000?"
"Aus Rechenzeit-Gründen. Das sind nur 3,7 % des Datensatzes – und genau
das ist unser größter Verbesserungshebel, dazu komme ich bei den
Herausforderungen." """,

4: """▶ TOKENIZER
"Wir benutzen einen Tokenizer auf Zeichen-Ebene: jedes einzelne Zeichen
ist ein Token. Vorteil: sehr einfach, und es gibt kein OOV-Problem –
es kann also kein unbekanntes Wort auftauchen."
• encode: Text → Zahlen-IDs   |   decode: IDs → Text

▶ GLEITENDES FENSTER (der Kern der Datenaufbereitung!)
"Wir schieben ein Fenster von 100 Zeichen über den Text, mit Schrittweite 50.
Das Ziel ist exakt der Input, um ein Zeichen nach rechts verschoben.
An jeder Position soll das Modell also das nächste Zeichen vorhersagen."

▶ auf das Beispiel rechts zeigen:
Input "K:D\\n|:A2B c2d"  →  Ziel ":D\\n|:A2B c2d|e"

💡 Das ist der Begriff, den man hören will: NEXT-TOKEN-PREDICTION.""",

5: """▶ DIE DREI BAUSTEINE (dem Pfeil-Diagramm folgen)

1) EMBEDDING (64 Dimensionen)
"Eine reine ID sagt nichts darüber aus, ob sich zwei Zeichen ähnlich
verhalten. Das Embedding macht daraus einen gelernten Vektor, sodass
musikalisch ähnliche Zeichen ähnliche Vektoren bekommen."

2) LSTM (2 Schichten à 256)
"Das LSTM verarbeitet die Sequenz Schritt für Schritt. Der Hidden State
ist sein Gedächtnis: er fasst alles zusammen, was bisher gelesen wurde."

3) LINEAR (→ 92 Logits) + Softmax
"Die Linear-Schicht gibt für jedes der 92 Zeichen einen Score aus,
Softmax macht daraus Wahrscheinlichkeiten."

📊 885.596 Parameter

❓ FALLS GEFRAGT "Warum LSTM und nicht ein einfaches RNN?"
"Ein einfaches RNN leidet beim Training am Vanishing-Gradient-Problem:
Das Lernsignal verblasst über lange Sequenzen. Das LSTM löst das mit
seinen Gates und einem Cell State, der die Sequenz durchläuft, ohne
bei jedem Schritt verfälscht zu werden." """,

6: """▶ SAGEN
"Trainiert wird mit CrossEntropy-Loss über alle Zeitschritte und dem
Adam-Optimizer."

📊 HYPERPARAMETER (stehen links in der Tabelle)
Batch 128 · Learning Rate 0,001 · 10 Epochen · Hidden 256 · 2 Schichten
· Sequenzlänge 100

▶ WANDB BETONEN (das wurde in der Aufgabe explizit verlangt!)
"Das komplette Training ist in Weights & Biases geloggt – ein Run pro
Experiment, beide im selben Projekt unter der Entity eidl-thm.
Dadurch kann ich LSTM und GRU direkt vergleichen."

▶ Pro Epoche geloggt: Loss, Top-1 und Top-5 – jeweils Train UND Validation.

❓ FALLS GEFRAGT "Warum Adam?"
"Adam passt die Schrittweite für jeden Parameter automatisch an –
das ist der Standard und läuft hier stabil." """,

7: """▶ LINKS: LOSS
"Der Loss fällt stetig – von 2,17 auf 1,06 im Training,
von 1,55 auf 1,08 in der Validierung."

▶ WICHTIG: Train- und Validierungskurve liegen fast aufeinander
"Es gibt also praktisch kein Overfitting."

▶ EHRLICHER PUNKT (bringt Pluspunkte!)
"Man sieht außerdem: die Validierungs-Loss fiel bei Epoche 10 immer noch.
Das Modell war also noch nicht auslerniert – wir hätten länger bzw. mit
mehr Daten trainieren können. Dazu komme ich beim Ausblick."

▶ RECHTS: ACCURACY
Top-1 steigt auf ca. 67 %, Top-5 auf ca. 93 %.

▶ Hinweis: "Die interaktiven Kurven aller Runs liegen in wandb." """,

8: """▶ DEFINITION ZUERST (genau wie im Übungsblatt!)
"Top-1 Accuracy ist der Anteil der Test-Sequenzen, bei denen das Modell
das nächste Zeichen als wahrscheinlichstes vorhersagt.
Top-5 ist der Anteil, bei dem das richtige Zeichen unter den fünf
wahrscheinlichsten liegt."

📊 ERGEBNIS
Top-1: 67,6 %   |   Top-5: 93,4 %

▶ DER SATZ, DER DIE ZAHL EINORDNET (nicht vergessen!)
"Zum Vergleich: bei 92 Zeichen läge reines Raten bei etwa 1,1 %.
Wir sind also um den Faktor 60 besser als der Zufall –
das Modell hat wirklich etwas gelernt."

❓ FALLS GEFRAGT "Wie haben Sie gesplittet?" (ehrliche Antwort = Pluspunkte)
"Ich habe die Fenster gemischt und dann 80/10/10 gesplittet. Da sich
benachbarte Fenster aber um 50 Zeichen überlappen, können zwei
überlappende Fenster in Train und Test landen – es gibt also eine leichte
Leakage, mein Testwert ist minimal optimistisch. Sauberer wäre: erst den
Text splitten, dann in jedem Teil die Fenster bauen. Der Effekt dürfte
hier aber klein sein, weil das Modell nicht overfittet – Train und
Validation liegen ja fast aufeinander, es hat also nichts auswendig
gelernt, das durchsickern könnte." """,

9: """▶ SAGEN
"Hier sehen Sie ein Stück, das mein LSTM selbst geschrieben hat.
Ich habe ihm nur den Header vorgegeben – also 6/8-Takt und D-Dur –
den Rest hat es Zeichen für Zeichen selbst erzeugt."

▶ AUF DIE STRUKTUR ZEIGEN
• M:6/8 = eine Jig, der typische irische Tanz; K:D = D-Dur
• Es setzt die Taktstriche | korrekt und hält die Tonart durch
• Es endet mit einem Schlussstrich ||

▶ DAS HIGHLIGHT (auf |1 ... :|2 ... zeigen!)
"Besonders interessant ist das hier: erste und zweite Endung.
Das ist eine typische Konvention irischer Volksmusik. Das habe ich
nirgendwo einprogrammiert – das Modell hat es allein aus den Daten
gelernt. Das zeigt, dass es wirklich langfristige Struktur erfasst."

▶ GENERIERUNG ERKLÄREN
"Generiert wird autoregressiv: Softmax über die Logits, dann ziehe ich
eine Stichprobe – statt einfach immer das wahrscheinlichste Zeichen zu
nehmen. Mit Argmax würde sich die Melodie sofort endlos wiederholen."

🔊 JETZT DIE WAV-DATEI ABSPIELEN (generierte_musik.wav) – der Aha-Moment!

⚠️ Falls jemand genau liest: "Das ist ein Ausschnitt, das vollständige
Stück liegt in beispiele.txt – dort wechselt es sogar die Tonart
nach D-Mixolydisch." """,

10: """▶ 1) ABLATION LSTM vs. GRU (2 Punkte)
"Ich habe exakt dieselbe Architektur nochmal trainiert und nur den
rekurrenten Layer getauscht."

📊 LSTM: 67,6 % / 93,4 %  –  885.596 Parameter
📊 GRU:  68,5 % / 93,6 %  –  671.580 Parameter  (24 % weniger!)

"Die GRU ist also minimal besser – bei 24 % weniger Parametern."

❓ FALLS GEFRAGT "Warum hat die GRU weniger Parameter?"
"Die GRU fusioniert Forget- und Input-Gate zu einem einzigen Update-Gate:
was sie neu aufnimmt, vergisst sie automatisch im gleichen Maß. Und sie
hat gar kein Output-Gate, weil sie Cell State und Hidden State zu einem
Zustand verschmilzt. Ergebnis: 3 Gewichtsmatrizen statt 4 – also genau
drei Viertel. Bei einer Aufgabe dieser Größe zahlt sich die zusätzliche
Flexibilität des LSTM offenbar nicht aus."

▶ 2) VALIDITY CHECK (2 Punkte)
"Top-1 sagt nichts darüber, ob ein GANZES Stück gültig ist. Deshalb prüfe
ich 6 Grammatikregeln der ABC-Notation: Kopfzeilen, Tonart, genügend
Taktstriche, echte Noten, keine leeren Takte, logischer Schluss.
Unsere generierten Stücke erfüllen 100 % der Regeln."

▶ 3) KREATIVITÄT (1 Punkt)
"Ein eigener Mini-Parser liest die Noten aus der ABC-Notation, jede Note
wird als Sinuswelle erzeugt (A4 = 440 Hz) und als WAV gespeichert –
ganz ohne Zusatzbibliothek." """,

11: """▶ EHRLICH BLEIBEN – das wird hier honoriert!

1) DATENMENGE vs. RECHENZEIT
"Wir haben nur 8.000 von 214.122 Stücken genutzt – also 3,7 %.
Das war ein bewusster Kompromiss für die Trainingszeit, ist aber
gleichzeitig unser größter Verbesserungshebel."

2) METRIK-DEFINITION
"Die Accuracy berechne ich laut Aufgabenstellung pro Sequenz auf dem
letzten Zeitschritt, den Loss aber über alle Zeitschritte – das gibt
dem Modell pro Sequenz viel mehr Lernsignal und beschleunigt das Lernen."

3) REPETITIVE MUSIK
"Mit Argmax wiederholte sich alles sofort – gelöst durch Sampling."

4) BEGRENZTER KONTEXT
"Das Modell hat beim Training nur 100 Zeichen Kontext gesehen. Sehr lange
Stücke verlieren daher manchmal die Struktur. Wir schneiden deshalb an
der ersten Leerzeile ab."

5) SPLIT / LEAKAGE (falls noch nicht auf Folie 8 gefragt – proaktiv sagen!)
"Und methodisch: meine gleitenden Fenster überlappen sich, und ich habe
sie vor dem Split gemischt. Streng genommen können sich Train und Test
dadurch leicht überschneiden – sauberer wäre, erst den Text zu splitten
und dann in jedem Teil die Fenster zu bauen." """,

12: """▶ ZUSAMMENFASSUNG
"Zusammengefasst: ein einfaches LSTM-Sprachmodell auf Zeichen-Ebene
reicht aus, um plausible irische Melodien zu erzeugen – 67,6 % Top-1
und 93,4 % Top-5 gegenüber 1,1 % beim Zufall. LSTM und GRU liegen
praktisch gleichauf, wobei die GRU mit deutlich weniger Parametern
auskommt."

▶ AUSBLICK (in dieser Reihenfolge – der erste Punkt ist der wichtigste!)
1. MEHR DATEN – wir nutzen nur 3,7 %. Da echte Vielfalt (andere Tonarten,
   andere Taktarten) mehr bringt als dieselben Stücke öfter zu sehen,
   ist das der stärkste Hebel.
2. Längerer Kontext / größeres Modell
3. Subword-Tokenizer (BPE) statt Zeichen-Ebene
4. Attention / Transformer (Vorlesung 10)

▶ ABSCHLUSS
"Vielen Dank für Ihre Aufmerksamkeit – gerne beantworten wir Ihre Fragen."

💡 Ruhig stehen bleiben, Fragen abwarten. Bei Unsicherheit:
"Das haben wir nicht untersucht, aber mein Ansatz wäre ..." """,
}


# ----------------------------------------------------------------- Ausführen
if not os.path.exists(PPTX):
    sys.exit("FEHLER: {} nicht gefunden. Bitte das Skript im selben Ordner "
             "wie die Präsentation ausführen.".format(PPTX))

# Sicherheitskopie, falls etwas schiefgeht
backup = "praesentation_ubung03_ohne_notizen.pptx.bak"
if not os.path.exists(backup):
    shutil.copy2(PPTX, backup)
    print("Sicherheitskopie angelegt: {}".format(backup))

prs = Presentation(PPTX)
print("Präsentation geöffnet: {} Folien".format(len(prs.slides._sldIdLst)))

for nummer, text in NOTIZEN.items():
    folie = prs.slides[nummer - 1]
    folie.notes_slide.notes_text_frame.text = text
    erste_zeile = text.strip().split("\n")[0]
    print("  Folie {:2d}: Notizen gesetzt  ({} Zeichen)".format(nummer, len(text)))

try:
    prs.save(PPTX)
    ziel = PPTX
except PermissionError:
    # Die Datei ist gerade in PowerPoint geöffnet -> in eine neue Datei speichern
    ziel = "praesentation_ubung03_MIT_NOTIZEN.pptx"
    prs.save(ziel)
    print("\nHINWEIS: {} ist gerade in PowerPoint geöffnet und konnte nicht "
          "überschrieben werden.".format(PPTX))
    print("Die Version mit Notizen wurde deshalb hier gespeichert: {}".format(ziel))

print("\nFertig! Notizen in {} gespeichert.".format(ziel))
print("In PowerPoint sichtbar über: Ansicht > Notizen  bzw. in der "
      "Referentenansicht während der Präsentation (Alt+F5 zum Testen).")
